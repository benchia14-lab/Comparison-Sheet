"""Local quote comparison dashboard. Uploaded files are processed in memory."""
from __future__ import annotations

import csv
import io
import math
import re
from difflib import SequenceMatcher
from pathlib import Path

import pymupdf as fitz
import openpyxl
from docx import Document
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024

HEADERS = {
    'description': ('description', 'item description', 'product', 'material', 'item', 'details', 'scope', 'work description'),
    'code': ('code', 'sku', 'product code', 'item code', 'part number', 'part no', 'ref', 'reference'),
    'qty': ('qty', 'quantity', 'units'),
    'unit': ('unit', 'uom', 'measure'),
    'unit_price': ('unit price', 'unit rate', 'rate', 'price each', 'each', 'unit cost', 'price'),
    'total': ('line total', 'amount', 'extended price', 'total', 'value'),
}
TOTAL_WORDS = re.compile(r'^(sub\s*total|total|gst|vat|tax|freight|delivery|shipping|discount|deposit|balance|page\s+\d+)', re.I)
CODE_RE = re.compile(r'\b(?=[A-Z0-9-]{4,}\b)(?=[A-Z0-9-]*\d)[A-Z][A-Z0-9-]*\b', re.I)


def money(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(float(value), 2)
    if value is None:
        return None
    s = str(value).strip().replace(',', '').replace('$', '').replace('£', '').replace('€', '')
    if s.startswith('(') and s.endswith(')'):
        s = '-' + s[1:-1]
    if re.fullmatch(r'-?\d+(?:\.\d+)?', s):
        return round(float(s), 2)
    return None


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def item(description, qty=None, unit_price=None, total=None, code='', unit='ea', source=''):
    description = clean(description)
    if not description or TOTAL_WORDS.match(description):
        return None
    qty, unit_price, total = money(qty), money(unit_price), money(total)
    if qty is None or qty <= 0:
        qty = 1.0 if unit_price is not None or total is not None else None
    if qty is None:
        return None
    if unit_price is None and total is not None:
        unit_price = round(total / qty, 2)
    if unit_price is None or unit_price < 0:
        return None
    inconsistent = total is not None and abs(total - qty * unit_price) > 0.02
    code = clean(code).upper()
    if not code:
        match = CODE_RE.search(description)
        code = match.group(0).upper() if match else ''
    return {'description': description, 'code': code, 'qty': qty, 'unit': clean(unit) or 'ea',
            'unit_price': unit_price, 'total': round(qty * unit_price, 2), 'source': source,
            'price_warning': inconsistent}


def header_map(cells):
    normalized = [re.sub(r'[^a-z0-9 ]', '', clean(c).lower()) for c in cells]
    result = {}
    for field, names in HEADERS.items():
        for i, cell in enumerate(normalized):
            if cell in names and field not in result:
                result[field] = i
    return result if 'description' in result and ('unit_price' in result or 'total' in result) else {}


def rows_to_items(rows, source_name):
    rows = list(rows)
    mapping, start = {}, 0
    for i, row in enumerate(rows[:30]):
        mapping = header_map(row)
        if mapping:
            start = i + 1
            break
    if not mapping:
        return []
    found = []
    for rowno, row in enumerate(rows[start:], start + 1):
        def val(key):
            idx = mapping.get(key)
            return row[idx] if idx is not None and idx < len(row) else None
        candidate = item(val('description'), val('qty'), val('unit_price'), val('total'),
                         val('code'), val('unit'), f'{source_name} row {rowno}')
        if candidate:
            found.append(candidate)
    return found


def parse_xlsx(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    # The supplied template contains 100 release sheets. Its first sheet is the
    # comparison source; subsequent tabs link back to it and are not new quotes.
    if 'Supplier Comparison Sheet' in wb.sheetnames:
        sheet = wb['Supplier Comparison Sheet']
        result = []
        for rowno, row in enumerate(sheet.iter_rows(min_row=7, max_col=4, values_only=True), 7):
            candidate = item(row[0], row[1], row[2], source=f'{sheet.title} row {rowno}')
            if candidate:
                result.append(candidate)
        return result, 'baseline'
    for sheet in wb:
        result = rows_to_items(sheet.iter_rows(values_only=True), sheet.title)
        if result:
            return result, 'quote'
    return [], 'quote'


def parse_csv(data):
    text = data.decode('utf-8-sig', errors='replace')
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return rows_to_items(csv.reader(io.StringIO(text), dialect), 'CSV')


def parse_text(text, source):
    found = []
    for no, raw in enumerate(text.splitlines(), 1):
        line = clean(raw)
        if not line or TOTAL_WORDS.match(line):
            continue
        # Common PDF/text quote layouts: description | qty | unit price | total,
        # or description [whitespace] qty unit price total.
        parts = [clean(p) for p in re.split(r'\t|\s{2,}|\s*\|\s*', raw) if clean(p)]
        if len(parts) >= 3:
            nums = [money(p) for p in parts]
            if nums[-1] is not None and nums[-2] is not None:
                if len(parts) >= 4 and nums[-3] is not None:
                    candidate = item(' '.join(parts[:-3]), nums[-3], nums[-2], nums[-1], source=f'{source} line {no}')
                else:
                    candidate = item(' '.join(parts[:-2]), nums[-2], total=nums[-1], source=f'{source} line {no}')
                if candidate and not candidate['price_warning']:
                    found.append(candidate)
    return found


def parse_pdf(data):
    with fitz.open(stream=data, filetype='pdf') as pdf:
        text = '\n'.join(page.get_text(sort=True) for page in pdf)
    return parse_pdf_text(text)


def parse_pdf_text(text):
    lines = text.splitlines()
    amount = r'([\d,]+(?:\.\d{1,2})?)'
    metadata = {'raw_text': text}
    if 'Product No.' in text and 'GST Inc Total' in text and 'BENTONS' in text.upper():
        for key, pattern_text in {
            'quote_number': r'Quote No\s+(\d+)',
            'quote_date': r'Quote Date\s+(\d{1,2}/\d{1,2}/\d{2,4})',
            'expiry_date': r'Expiry Date\s+(\d{1,2}/\d{1,2}/\d{2,4})',
        }.items():
            found_meta = re.search(pattern_text, text, re.I)
            if found_meta: metadata[key] = found_meta.group(1)
        pattern = re.compile(r'^\s*(BB\d+)\s+(.+?)\s+(EA|EACH|PC|PCS)\s+' + amount + r'\s+' + amount + r'\s+' + amount + r'\s+' + amount + r'\s*$', re.I)
        found = []
        for no, raw in enumerate(lines, 1):
            match = pattern.match(raw)
            if match:
                code, description, unit, qty, ex, inc, total = match.groups()
                entry = item(description, qty, ex, source=f'PDF line {no}', code=code, unit=unit)
                if entry:
                    details = []
                    for continuation in lines[no:]:
                        value = clean(continuation)
                        if re.match(r'^\$[\d,.]+$', value) or re.match(r'^BB\d+', value) or value.startswith('*Quotes'):
                            break
                        if value and value not in ('SNK1', 'TA1'):
                            details.append(value)
                    entry['full_description'] = ' '.join([description, *details])
                    entry['inc_gst_unit'] = money(inc)
                    entry['inc_gst_total'] = money(total)
                    found.append(entry)
        for line in lines:
            gst = re.search(r'\bGST\s+(?!Ex|Inc)' + r'.*?\$' + amount, line, re.I)
            total = re.search(r'GST Inc Total\s+\$' + amount, line, re.I)
            if gst: metadata['tax'] = money(gst.group(1))
            if total: metadata['quoted_total'] = money(total.group(1))
        metadata.update({'name': "Benton's", 'price_basis': 'ex GST'})
        return found, metadata
    if 'Reece Australia Pty Ltd' in text and 'Document Total' in text:
        for key, pattern_text in {
            'quote_number': r'Quote No\.\s+(\d+)',
            'quote_date': r'\bDate\s+(\d{1,2}/\d{1,2}/\d{2,4})',
            'expiry_date': r'Quote Expiry Date:\s*(\d{1,2}/\d{1,2}/\d{2,4})',
        }.items():
            found_meta = re.search(pattern_text, text, re.I)
            if found_meta: metadata[key] = found_meta.group(1)
        product = re.compile(r'^\s*\d+\s+\d+\s+([A-Za-z0-9]+)\s+-?\s*(.+?)\s+' + amount + r'\s+' + amount + r'\s+(EA|EACH|PC|PCS)\s+' + amount + r'\s+' + amount + r'\s+' + amount + r'\s*$', re.I)
        freight = re.compile(r'Delivery Fee -.+?\s+' + amount + r'\s+' + amount + r'\s+EA\s+' + amount + r'\s+' + amount + r'\s+' + amount, re.I)
        found = []
        for no, raw in enumerate(lines, 1):
            match = product.match(raw)
            if match:
                code, description, qty, price, unit, net, gst, gross = match.groups()
                entry = item(description, qty, price, net, code, unit, f'PDF line {no}')
                if entry:
                    details = []
                    for continuation in lines[no:]:
                        value = clean(continuation)
                        if value.startswith('*') or re.match(r'^\d+\s+', value) or 'Delivery Fee' in value:
                            break
                        if value:
                            details.append(value)
                    entry['full_description'] = ' '.join([description, *details])
                    entry['inc_gst_total'] = money(gross)
                    found.append(entry)
            shipping = freight.search(raw)
            if shipping:
                metadata['freight'] = money(shipping.group(3))
            total = re.search(r'Document Total\s+' + amount + r'\s+' + amount + r'\s+' + amount, raw, re.I)
            if total:
                metadata['tax'] = money(total.group(2))
                metadata['quoted_total'] = money(total.group(3))
        metadata.update({'name': 'Reece', 'price_basis': 'ex GST'})
        return found, metadata
    return parse_text(text, 'PDF'), metadata


def parse_document(filename, data):
    suffix = Path(filename).suffix.lower()
    if suffix == '.xlsx':
        entries, kind = parse_xlsx(data)
        return entries, kind, {}
    if suffix == '.csv':
        return parse_csv(data), 'quote', {'raw_text': data.decode('utf-8-sig', errors='replace')}
    if suffix == '.pdf':
        entries, metadata = parse_pdf(data)
        return entries, 'quote', metadata
    if suffix == '.docx':
        doc = Document(io.BytesIO(data))
        rows = []
        for table in doc.tables:
            rows += [[cell.text for cell in row.cells] for row in table.rows]
        tabular = rows_to_items(rows, 'Word table')
        raw_text = '\n'.join([*(p.text for p in doc.paragraphs), *("\t".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)])
        return tabular or parse_text(raw_text, 'Word'), 'quote', {'raw_text': raw_text}
    if suffix == '.txt':
        raw_text = data.decode('utf-8-sig', errors='replace')
        return parse_text(raw_text, 'Text'), 'quote', {'raw_text': raw_text}
    raise ValueError('Supported files: XLSX, CSV, PDF, DOCX, TXT')


def parse_file(filename, data):
    entries, kind, _ = parse_document(filename, data)
    return entries, kind


def norm(s):
    return re.sub(r'[^a-z0-9]+', ' ', clean(s).lower()).strip()


def match_score(a, b):
    if a['code'] and b['code'] and norm(a['code']) == norm(b['code']):
        return 1.0
    x, y = norm(a['description']), norm(b['description'])
    if not x or not y:
        return 0.0
    stop = {'the', 'with', 'and', 'for', 'ea', 'each', 'sink', 'bowl', 'w', 'xl'}
    words_x, words_y = set(x.split()) - stop, set(y.split()) - stop
    shared = words_x & words_y
    if len(shared) < 2:
        return 0.0
    overlap = len(shared) / max(1, min(len(words_x), len(words_y)))
    return 0.65 * overlap + 0.35 * SequenceMatcher(None, x, y).ratio()


def build_groups(quotes, baseline):
    groups = []
    for entry in baseline:
        groups.append({'description': entry['description'], 'code': entry['code'], 'qty': entry['qty'],
                       'unit': entry['unit'], 'baseline_price': entry['unit_price'], 'matches': {},
                       'confidence': 'baseline'})
    for qi, quote in enumerate(quotes):
        for entry in quote['items']:
            candidates = [(match_score(entry, group), gi) for gi, group in enumerate(groups)
                          if qi not in group['matches']]
            score, gi = max(candidates, default=(0, -1))
            if score >= 0.70:
                group = groups[gi]
                group['matches'][qi] = entry
                if score < 0.93:
                    group['confidence'] = 'review'
                elif group['confidence'] in ('new', 'unmatched'):
                    group['confidence'] = 'matched'
            else:
                groups.append({'description': entry['description'], 'code': entry['code'], 'qty': entry['qty'],
                               'unit': entry['unit'], 'baseline_price': None, 'matches': {qi: entry},
                               'confidence': 'unmatched' if qi else 'new'})
    return groups


def compare(quotes, groups):
    names = [q['name'] for q in quotes]
    totals = [0.0] * len(quotes)
    covered = [0] * len(quotes)
    rows = []
    for group in groups:
        qty = money(group.get('qty')) or 1
        prices = []
        for qi in range(len(quotes)):
            match = group.get('matches', {}).get(str(qi)) or group.get('matches', {}).get(qi)
            if match:
                try:
                    unit_price = float(match.get('unit_price'))
                    prices.append(round(unit_price, 6) if math.isfinite(unit_price) and unit_price >= 0 else None)
                except (TypeError, ValueError):
                    prices.append(None)
            else:
                prices.append(None)
        available = [p for p in prices if p is not None]
        best = min(available) if available else None
        values = [round(qty * p, 2) if p is not None else None for p in prices]
        for qi, value in enumerate(values):
            if value is not None:
                totals[qi] += value
                covered[qi] += 1
        rows.append({**group, 'prices': prices, 'values': values, 'best': best,
                     'winners': [names[i] for i, p in enumerate(prices) if p is not None and p == best],
                     'spread': round((max(available) - best) * qty, 2) if len(available) > 1 else None,
                     'baseline_total': round(qty * group['baseline_price'], 2) if group.get('baseline_price') is not None else None})
    complete = [covered[i] == len(rows) for i in range(len(quotes))]
    landed = []
    for i, quote in enumerate(quotes):
        freight = money(quote.get('freight')) or 0
        tax = money(quote.get('tax')) or 0
        discount = money(quote.get('discount')) or 0
        landed.append(round(totals[i] + freight + tax - discount, 2))
    return {'rows': rows, 'names': names, 'totals': [round(x, 2) for x in totals],
            'landed': landed, 'covered': covered, 'complete': complete, 'row_count': len(rows)}


@app.get('/')
def home():
    return render_template('index.html')


@app.post('/api/import')
def import_files():
    result = []
    for f in request.files.getlist('files'):
        try:
            entries, kind, metadata = parse_document(f.filename or '', f.read())
            result.append({'name': metadata.get('name') or Path(f.filename).stem, 'kind': kind, 'items': entries,
                           **metadata,
                           'warning': None if entries else 'No priced line items found. Check the source layout or enter rows manually.'})
        except Exception as exc:
            result.append({'name': f.filename, 'kind': 'error', 'items': [], 'warning': str(exc)})
    return jsonify(result)


@app.post('/api/compare')
def compare_api():
    data = request.get_json(force=True)
    quotes = data.get('quotes', [])
    groups = data.get('groups')
    if groups is None:
        groups = build_groups(quotes, data.get('baseline', []))
    return jsonify(compare(quotes, groups))


@app.post('/api/export')
def export():
    data = request.get_json(force=True)
    quotes = data.get('quotes', [])
    report = compare(quotes, data.get('groups', []))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Comparison'
    ws.append(['QUOTE COMPARISON'])
    ws.append(['Project', clean(data.get('project'))])
    ws.append(['Prices are compared at the required quantity; missing lines are excluded from quote totals.'])
    header = ['Code', 'Description', 'Required Qty', 'Unit', 'Baseline Unit', 'Baseline Total']
    for name in report['names']:
        header += [f'{name} Quoted Code', f'{name} Quoted Description', f'{name} Quoted Qty',
                   f'{name} Quoted Unit', f'{name} Unit Price', f'{name} At Required Qty', f'{name} Source', f'{name} Price Verified']
    header += ['Lowest Unit', 'Lowest Supplier', 'Price Spread at Required Qty', 'Review Flag']
    ws.append(header)
    for row in report['rows']:
        values = [row.get('code'), row.get('description'), row.get('qty'), row.get('unit'),
                  row.get('baseline_price'), row.get('baseline_total')]
        for qi in range(len(quotes)):
            match = row.get('matches', {}).get(str(qi)) or row.get('matches', {}).get(qi)
            values += [match.get('code') if match else None, (match.get('full_description') or match.get('description')) if match else None,
                       match.get('qty') if match else None, match.get('unit') if match else None,
                       row['prices'][qi], row['values'][qi], match.get('source') if match else 'MISSING',
                       ('Yes' if match.get('price_verified') else 'No') if match else None]
        values += [row['best'], ', '.join(row['winners']), row['spread'], row.get('confidence')]
        ws.append(values)
    ws.append([])
    ws.append(['QUOTE TOTALS', '', '', '', '', ''] + [x for i in range(len(quotes)) for x in ('', '', '', '', '', report['totals'][i], '', '')])
    ws.append(['LINES COVERED', '', '', '', '', ''] + [x for i in range(len(quotes)) for x in ('', '', '', '', '', f"{report['covered'][i]}/{report['row_count']}", '', '')])
    for cell in ws[4]:
        cell.fill = PatternFill('solid', fgColor='203B38')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'C5'
    ws.auto_filter.ref = f'A4:{get_column_letter(ws.max_column)}{4 + len(report["rows"])}'
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 52
    for col in range(3, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 17
    for col, title in enumerate(header, 1):
        if 'Description' in title:
            ws.column_dimensions[get_column_letter(col)].width = 42
        elif 'Source' in title:
            ws.column_dimensions[get_column_letter(col)].width = 24
    for row in ws.iter_rows(min_row=5, max_row=4 + len(report['rows'])):
        for cell in row:
            if isinstance(cell.value, (int, float)) and cell.column != 3:
                cell.number_format = '#,##0.00;[Red](#,##0.00)'
    detail = wb.create_sheet('Source lines')
    detail.append(['Quote', 'Code', 'Description', 'Quoted Qty', 'Unit', 'Quoted Unit Price', 'Quoted Line Total', 'Source', 'Price Verified', 'Verified At (UTC)'])
    for quote in quotes:
        for entry in quote['items']:
            detail.append([quote['name'], entry.get('code'), entry.get('full_description') or entry.get('description'), entry.get('qty'),
                           entry.get('unit'), entry.get('unit_price'), entry.get('total'), entry.get('source'),
                           'Yes' if entry.get('price_verified') else 'No', entry.get('price_verified_at')])
    detail.freeze_panes = 'C2'
    detail.auto_filter.ref = detail.dimensions
    detail.column_dimensions['C'].width = 52
    detail.column_dimensions['H'].width = 30
    detail.column_dimensions['I'].width = 17
    detail.column_dimensions['J'].width = 25
    summary = wb.create_sheet('Financial summary', 0)
    summary.append(['Project', clean(data.get('project'))])
    summary.append(['Quoted line costs use required quantities. Enter tax, freight and discounts after checking quote terms.'])
    summary.append(['Supplier', 'Price basis', 'Lines covered', 'Complete?', 'Line cost', 'Freight', 'Tax / GST', 'Discount', 'Total evaluated', 'Quoted total', 'Variance to quote', 'Gap to lowest complete'])
    complete_totals = [report['landed'][i] for i, yes in enumerate(report['complete']) if yes]
    lowest = min(complete_totals) if complete_totals else None
    for i, quote in enumerate(quotes):
        quoted_total = money(quote.get('quoted_total'))
        summary.append([quote['name'], quote.get('price_basis', 'unspecified'),
                        f"{report['covered'][i]}/{report['row_count']}",
                        'Yes' if report['complete'][i] else 'No', report['totals'][i],
                        money(quote.get('freight')) or 0, money(quote.get('tax')) or 0,
                        money(quote.get('discount')) or 0, report['landed'][i], quoted_total,
                        round(report['landed'][i] - quoted_total, 2) if quoted_total is not None else None,
                        round(report['landed'][i] - lowest, 2) if report['complete'][i] and lowest is not None else None])
    for cell in summary[3]:
        cell.fill = PatternFill('solid', fgColor='203B38')
        cell.font = Font(color='FFFFFF', bold=True)
    summary.freeze_panes = 'B4'
    summary.column_dimensions['A'].width = 25
    for col in range(2, 13):
        summary.column_dimensions[get_column_letter(col)].width = 19
    # Prevent a supplier-controlled description or filename from becoming a
    # spreadsheet formula when the exported workbook is opened.
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                    cell.value = "'" + cell.value
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='quote-comparison.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=8000, debug=False)
